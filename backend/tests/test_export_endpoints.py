# Copyright (C) 2026 CenkorMES Project
# SPDX-License-Identifier: AGPL-3.0
"""列表导出端点测试：库存、客户对账单。

这两个导出原先只有前端调用、后端没挂路由（点「导出 Excel」必 404）。
导出是同步流式 xlsx（本机没有 cenkormes 的 celery worker，异步 export-job 链路等不到结果），
所以断言直接落在响应字节上。
"""
import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.api.admin.finance.router import export_statements_api
from app.api.admin.warehouse.router import export_stocks_api
from app.crud.warehouse import adjust_stock, create_warehouse
from app.main import app
from app.models.customer import Customer
from app.models.finance import Statement
from app.models.sku import Sku


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _body(resp) -> bytes:
    chunks = [
        chunk if isinstance(chunk, bytes) else chunk.encode()
        async for chunk in resp.body_iterator
    ]
    return b"".join(chunks)


def _sheet(resp) -> list[list]:
    assert resp.media_type == XLSX_MIME
    assert "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(asyncio.run(_body(resp))))
    return [list(row) for row in wb.active.iter_rows(values_only=True)]


def _mounted_paths() -> set:
    out: set = set()

    def walk(routes):
        for r in routes:
            p = getattr(r, "path", None)
            if p:
                out.add(p)
            if getattr(r, "routes", None):
                walk(r.routes)

    walk(app.routes)
    return out


# ── 库存导出 ──

def test_stocks_export_route_mounted():
    assert "/api/admin/warehouse/stocks/export" in _mounted_paths()


def test_statements_export_route_mounted():
    assert "/api/admin/finance/statements/export" in _mounted_paths()


def test_stocks_export_streams_rows(session: Session, test_user, sku: Sku):
    wh = create_warehouse(session, code="WH01", name="一号仓")
    adjust_stock(session, wh.id, sku.id, 12, "manual")
    rows = _sheet(export_stocks_api(
        warehouse_id=None, item_type=None, db=session, user=test_user
    ))
    assert rows[0] == ["仓库", "编码", "名称", "数量", "更新时间"]
    assert rows[1][0] == "一号仓"
    assert rows[1][1] == "SKU001"
    assert rows[1][3] == 12


def test_stocks_export_honours_filters(session: Session, test_user, sku: Sku):
    a = create_warehouse(session, code="WA", name="A仓")
    b = create_warehouse(session, code="WB", name="B仓")
    adjust_stock(session, a.id, sku.id, 5, "manual")
    adjust_stock(session, b.id, sku.id, 7, "manual")

    only_a = _sheet(export_stocks_api(
        warehouse_id=a.id, item_type=None, db=session, user=test_user
    ))
    assert [r[0] for r in only_a[1:]] == ["A仓"]

    # SKU001 不是 MAT- 前缀 → 只看物料时应当为空表（只留表头）
    empty = _sheet(export_stocks_api(
        warehouse_id=None, item_type="material", db=session, user=test_user
    ))
    assert len(empty) == 1


# ── 客户对账单导出 ──

def _statement(session: Session, code: str, customer: Customer, total: str, status: str) -> Statement:
    s = Statement(
        customer_id=customer.id,
        code=code,
        period_start=date(2026, 1, 1),
        period_end=date(2026, 1, 31),
        total_amount=Decimal(total),
        paid_amount=Decimal("0"),
        status=status,
        due_date=date(2026, 2, 28),
        remark=None,
    )
    session.add(s)
    session.flush()
    return s


def test_statements_export_columns_and_amounts(session: Session, test_user, customer: Customer):
    _statement(session, "ST-1", customer, "1000.0000", "draft")
    rows = _sheet(export_statements_api(
        customer_id=None, status=None, db=session, user=test_user
    ))
    assert rows[0] == [
        "对账单号", "客户", "期间起", "期间止", "应收金额", "已核销", "未收余额", "到期日", "状态", "备注",
    ]
    row = rows[1]
    assert row[0] == "ST-1"
    assert row[1] == customer.name
    assert row[4] == 1000 and row[5] == 0 and row[6] == 1000
    assert row[7] == "2026-02-28"
    assert row[8] == "draft"


def test_statements_export_not_truncated_by_list_pagination(session: Session, test_user, customer: Customer):
    """列表接口默认 limit=50；导出必须给全，否则用户以为筛出来就这么多。"""
    for i in range(60):
        _statement(session, f"ST-{i:03d}", customer, "1.0000", "draft")
    rows = _sheet(export_statements_api(
        customer_id=None, status=None, db=session, user=test_user
    ))
    assert len(rows) == 61  # 1 表头 + 60 行


def test_statements_export_honours_filters(session: Session, test_user, customer: Customer, sku: Sku):
    other = Customer(code="C002", name="另一客户", is_active=True)
    session.add(other)
    session.flush()

    _statement(session, "ST-A", customer, "1.0000", "draft")
    _statement(session, "ST-B", customer, "2.0000", "paid")
    _statement(session, "ST-C", other, "3.0000", "draft")

    by_customer = _sheet(export_statements_api(
        customer_id=customer.id, status=None, db=session, user=test_user
    ))
    assert sorted(r[0] for r in by_customer[1:]) == ["ST-A", "ST-B"]

    by_status = _sheet(export_statements_api(
        customer_id=None, status="draft", db=session, user=test_user
    ))
    assert sorted(r[0] for r in by_status[1:]) == ["ST-A", "ST-C"]


def test_statements_export_empty_is_valid_workbook(session: Session, test_user, customer: Customer):
    rows = _sheet(export_statements_api(
        customer_id=customer.id, status="paid", db=session, user=test_user
    ))
    assert len(rows) == 1

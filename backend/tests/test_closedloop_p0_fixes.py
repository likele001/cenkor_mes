# Copyright (C) 2026 CenkorMES Project
# SPDX-License-Identifier: AGPL-3.0
"""闭环体检第一批（P0 止血）回归测试

全部走隔离 SQLite + 直接调用 handler，不联网、不碰生产库。
直接调用 handler 时 FastAPI 不会把 Query(...) 默认值还原成 None，
所以未筛选的查询参数必须显式给值。
"""
from __future__ import annotations

import ast
import asyncio
import importlib
import inspect
import pathlib
from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.customer import Customer
from app.models.print_template import PrintTemplate
from app.models.report import Report
from app.models.sku import Sku
from app.models.task import Task
from app.models.department import Department
from app.models.user import User, user_roles
from app.models.work_order import WorkOrder


# ────────────────────────── 守门员：模块内导入必须能解析 ──────────────────────────

OPTIONAL_MODULES = {
    # 独立版没有 AI 员工模块，飞书回调里用 try 兜底降级为回执文案
    "app.services.ai_employee.im_dispatch",
}


def _iter_from_imports(node: ast.AST):
    for child in ast.iter_child_nodes(node):
        if isinstance(child, ast.ImportFrom):
            yield child
        yield from _iter_from_imports(child)


def _submodule_exists(pkg, name: str) -> bool:
    return any(
        (pathlib.Path(p) / f"{name}.py").exists() or (pathlib.Path(p) / name).is_dir()
        for p in (getattr(pkg, "__path__", []) or [])
    )


def test_all_intra_app_imports_resolve():
    """app/ 里每一条 `from app.x import y` 都必须真的存在。

    闭环体检发现的最大一类问题就是这个：service 文件被删/改名后调用点仍写旧名字，
    只有等用户点开那个页面才 500。
    """
    root = pathlib.Path("app")
    problems: list[str] = []
    for f in sorted(root.rglob("*.py")):
        try:
            tree = ast.parse(f.read_text())
        except SyntaxError as e:
            problems.append(f"{f}: SyntaxError {e}")
            continue
        for node in _iter_from_imports(tree):
            if not (node.module or "").startswith("app.") or node.module in OPTIONAL_MODULES:
                continue
            try:
                mod = importlib.import_module(node.module)
            except Exception as e:  # noqa: BLE001
                problems.append(f"{f}:{node.lineno} from {node.module} -> {type(e).__name__}: {e}")
                continue
            for a in node.names:
                if hasattr(mod, a.name):
                    continue
                if hasattr(mod, "__path__") and _submodule_exists(mod, a.name):
                    continue  # from package import submodule
                problems.append(f"{f}:{node.lineno} from {node.module} import {a.name} -> MISSING")
    assert not problems, "\n".join(problems)


def test_approval_flow_resolver_call_sites_have_right_arity():
    """件次审核曾因为少传 db 给 is_terminal_status/format_step_label 而必崩。"""
    from app.services import approval_flow_resolver as r

    for fn in (r.is_terminal_status, r.format_step_label):
        params = list(inspect.signature(fn).parameters)
        assert params[0] == "db", f"{fn.__name__} 首参必须是 db"

    src = importlib.import_module("app.api.admin.production.report_units").__file__ or ""
    source = pathlib.Path(src).read_text()
    assert "is_terminal_status(db," in source
    assert "format_step_label(db," in source
    assert "is_terminal_status(new_status)" not in source
    assert "format_step_label(next_step," not in source


# ────────────────────────── 1. 入库单列表返回 items ──────────────────────────

def _make_warehouse(session: Session):
    from app.models.warehouse import Warehouse

    wh = session.query(Warehouse).filter(Warehouse.name == "测试仓").first()
    if not wh:
        wh = Warehouse(code="WH-TEST", name="测试仓", is_active=True)
        session.add(wh)
        session.flush()
    return wh


def test_entries_list_returns_items_key(session: Session, test_user: User, sku: Sku):
    from app.api.admin.warehouse.warehouse_entries import list_api
    from app.crud.warehouse_entry import create_entry

    res = list_api(warehouse_id=None, source_type=None, status=None, offset=0, limit=50, db=session)
    assert isinstance(res["data"], dict), "前端读 res.items，列表接口不能返回裸数组"
    assert "items" in res["data"]

    wh = _make_warehouse(session)
    from app.models.material import Material

    mat = session.query(Material).filter(Material.code == "MAT-TEST-1").first()
    if not mat:
        mat = Material(code="MAT-TEST-1", name="测试物料", sku_id=sku.id, is_active=True)
        session.add(mat)
        session.flush()
    entry = create_entry(
        db=session,
        code="WE-TEST-1",
        source_type="other",
        warehouse_id=wh.id,
        items=[{"material_id": mat.id, "sku_id": sku.id, "qty": 1}],
        created_by=test_user.id,
    )
    session.commit()

    res2 = list_api(warehouse_id=None, source_type=None, status=None, offset=0, limit=50, db=session)
    assert any(x["code"] == entry.code for x in res2["data"]["items"])


# ────────────────────────── 2. 工单导出不再引用 x.code ──────────────────────────

def _sheet_rows(resp) -> list[list]:
    async def _body():
        return b"".join([c if isinstance(c, bytes) else c.encode() async for c in resp.body_iterator])

    wb = load_workbook(BytesIO(asyncio.run(_body())))
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


def test_work_order_export_returns_readable_rows(session: Session, work_order: WorkOrder, test_user: User):
    from app.api.admin.production.work_orders import export_api

    resp = export_api(order_id=None, status=None, db=session, user=test_user)
    rows = _sheet_rows(resp)
    assert rows[0][0] == "工单号"
    assert rows[1][0] == work_order.id
    assert work_order.order.code in rows[1]


def test_work_order_export_honours_status_filter(session: Session, work_order: WorkOrder, test_user: User):
    from app.api.admin.production.work_orders import export_api

    rows = _sheet_rows(export_api(order_id=None, status="nope", db=session, user=test_user))
    assert len(rows) == 1, "导出必须与列表用同一套筛选"


# ────────────────────────── 3. 客户打印 / 打印 PDF ──────────────────────────

def test_customer_print_without_template_gives_actionable_400(session: Session, customer: Customer, test_user: User):
    from app.api.admin.production.customers import print_api

    with pytest.raises(HTTPException) as exc:
        print_api(customer_id=customer.id, template_id=None, template_code="customer_card_absent", db=session, user=test_user)
    assert exc.value.status_code == 400
    assert "打印模板" in exc.value.detail


def test_customer_print_renders_template(session: Session, customer: Customer, test_user: User):
    from app.api.admin.production.customers import print_api

    session.add(PrintTemplate(
        code="customer_card", name="客户名片", template_type="html",
        content="<div>{{ customer.name }} / {{ customer.code }} / {{ customer.owner_name }}</div>",
        is_active=True))
    session.commit()

    res = print_api(customer_id=customer.id, template_id=None, template_code="customer_card", db=session, user=test_user)
    html = res["data"]["html"]
    assert customer.name in html and customer.code in html
    assert res["data"]["customer_id"] == customer.id


def test_customer_print_routes_are_mounted():
    from app.main import app

    paths = {getattr(r, "path", "") for r in app.routes}
    assert "/api/admin/production/customers/{customer_id}/print" in paths
    assert "/api/admin/production/customers/{customer_id}/print-pdf" in paths


# ────────────────────────── 4. 计划交期预测 ──────────────────────────

@pytest.fixture
def plan(session: Session, order_item: tuple):
    from app.models.production_plan import ProductionPlan

    order, _item = order_item
    p = ProductionPlan(
        order_id=order.id, code="PP-TEST-1", status="planned",
        start_date=date.today() - timedelta(days=1),
        end_date=date.today() + timedelta(days=9),
        work_days=10,
    )
    session.add(p)
    session.commit()
    return p


def _approve_report(session: Session, task: Task, user: User, qty: int, days_ago: int = 1) -> Report:
    r = Report(task_id=task.id, report_user_id=user.id, good_qty=qty, bad_qty=0, status="qc_approved")
    session.add(r)
    session.flush()
    # created_at 是 server_default，回填到 N 天前才能落进近 7 天产能窗口
    session.execute(
        Report.__table__.update()
        .where(Report.id == r.id)
        .values(created_at=datetime.now() - timedelta(days=days_ago))
    )
    session.commit()
    return r


def test_plan_forecast_reports_real_numbers(session: Session, plan, task: Task, test_user: User):
    from app.services.production_forecast import build_plan_forecast

    _approve_report(session, task, test_user, qty=70, days_ago=1)
    out = build_plan_forecast(db=session, plan_id=plan.id)

    assert out["plan_id"] == plan.id
    assert out["remaining_tasks"] == 1
    assert out["remaining_qty"] == 100
    assert out["avg_daily_output_7d"] == pytest.approx(10.0, abs=0.01)
    assert out["days_needed"] == pytest.approx(10.0, abs=0.1)
    assert isinstance(out["kitting_ok"], bool)
    assert out["due_risk"] in ("low", "medium", "high", "overdue", "unknown")


def test_plan_forecast_ignores_unapproved_reports(session: Session, plan, task: Task, test_user: User):
    from app.services.production_forecast import build_plan_forecast

    session.add(Report(task_id=task.id, report_user_id=test_user.id, good_qty=500, bad_qty=0, status="submitted"))
    session.commit()
    out = build_plan_forecast(db=session, plan_id=plan.id)
    assert out["avg_daily_output_7d"] == 0
    assert out["days_needed"] is None


def test_plan_forecast_flags_overdue(session: Session, plan, order_item):
    from app.services.production_forecast import build_plan_forecast

    order, _item = order_item
    order.due_date = date.today() - timedelta(days=3)
    session.commit()

    out = build_plan_forecast(db=session, plan_id=plan.id)
    assert out["days_left"] == -3
    assert out["due_risk"] == "overdue"


def test_plan_forecast_missing_plan_raises(session: Session):
    from app.services.production_forecast import build_plan_forecast

    with pytest.raises(ValueError):
        build_plan_forecast(db=session, plan_id=999999)


# ────────────────────────── 5. APS 策略对比 ──────────────────────────

def test_aps_strategies_are_comparable(session: Session, plan, task: Task, test_user: User):
    from app.services.aps_strategy_analysis import analyze_aps_strategies

    out = analyze_aps_strategies(db=session, plan_id=plan.id, user_id=test_user.id)
    keys = {s["key"] for s in out["strategies"]}
    assert {"backward", "forward", "optimized"} <= keys
    scores = [s["score"] for s in out["strategies"]]
    assert scores == sorted(scores, reverse=True), "策略必须按得分降序，推荐位才可信"
    assert out["recommended"] in keys
    assert out["forecast"]["plan_id"] == plan.id
    for s in out["strategies"]:
        assert s["pros"] or s["cons"], f"{s['key']} 没有任何理由，评分不可解释"


def test_aps_route_defined_only_once():
    from app.api.admin.production.plans import router

    paths = [getattr(r, "path", "") for r in router.routes]
    for p in ("/plans/{plan_id}/forecast", "/plans/{plan_id}/aps-strategy"):
        assert paths.count(p) == 1, f"{p} 重复定义 {paths.count(p)} 次"


# ────────────────────────── 6. 计划自动派工（service 化 + 单租户签名） ──────────────────────────

@pytest.fixture
def worker(session: Session, employee_role, department: Department) -> User:
    u = User(
        username="worker01",
        password_hash="x",
        full_name="张员工",
        is_active=True,
        department_id=department.id,
    )
    session.add(u)
    session.flush()
    session.execute(user_roles.insert().values(user_id=u.id, role_id=employee_role.id))
    session.flush()
    return u


@pytest.fixture
def released_plan(session: Session, plan):
    plan.status = "in_progress"
    session.commit()
    return plan


def test_auto_dispatch_signature_has_no_tenant_id():
    """自动化编排曾按旧的 tenant_id 签名调用，单租户改造后这条链路必崩。"""
    from app.services.plan_auto_dispatch import execute_auto_dispatch

    params = inspect.signature(execute_auto_dispatch).parameters
    assert "tenant_id" not in params
    assert {"plan_id", "user_id", "payload"} <= set(params)

    src = pathlib.Path(inspect.getfile(execute_auto_dispatch)).read_text()
    assert "tenant_id=" not in src, "派工 service 里不应再向 crud 传 tenant_id"


def test_auto_dispatch_refuses_unreleased_plan(session: Session, plan):
    from app.services.plan_auto_dispatch import execute_auto_dispatch
    from app.schemas.production_plan import AutoDispatchIn

    with pytest.raises(ValueError, match="尚未"):
        execute_auto_dispatch(
            session, plan_id=plan.id, user_id=1, payload=AutoDispatchIn())


def test_auto_dispatch_raises_when_no_workers(session: Session, released_plan):
    from app.services.plan_auto_dispatch import execute_auto_dispatch
    from app.schemas.production_plan import AutoDispatchIn

    with pytest.raises(ValueError, match="无可用员工"):
        execute_auto_dispatch(
            session,
            plan_id=released_plan.id,
            user_id=1,
            payload=AutoDispatchIn(),
        )


def test_auto_dispatch_assigns_tasks_and_reports_load(
    session: Session, task: Task, worker: User, released_plan
):
    from app.services.plan_auto_dispatch import execute_auto_dispatch
    from app.schemas.production_plan import AutoDispatchIn

    out = execute_auto_dispatch(
        session,
        plan_id=released_plan.id,
        user_id=worker.id,
        payload=AutoDispatchIn(user_ids=[worker.id]),
    )
    assert out["task_count"] == 1
    assert out["assigned_count"] == 1
    assert {"unit", "users", "workshops", "overloads", "release"} <= set(out)
    assert out["users"][0]["user_id"] == worker.id
    assert out["users"][0]["total_minutes"] > 0

    from app.crud.task_assignment import list_assignments_for_task

    assert list_assignments_for_task(session, task.id), "派工必须真的落成 TaskAssignment"


def test_auto_dispatch_route_delegates_to_service():
    """plans API 里不能再留第二份派工实现。"""
    import app.api.admin.production.plans as m

    src = pathlib.Path(inspect.getfile(m)).read_text()
    assert "from app.services.plan_auto_dispatch import execute_auto_dispatch" in src
    assert "ws_load" not in src, "路由里不应再出现派工算法本身"


# ────────────────────────── 7. 自动化编排签名 ──────────────────────────

def test_automation_helpers_have_no_tenant_id_params():
    import app.services.production_automation as m

    for name, fn in inspect.getmembers(m, inspect.isfunction):
        if not name.startswith(("_", "run_", "precheck_", "auto_", "apply_", "maybe_", "enqueue_", "log_")):
            continue
        if fn.__module__ != m.__name__:
            continue
        assert "tenant_id" not in inspect.signature(fn).parameters, f"{name} 仍带 tenant_id 形参"


def test_maybe_trigger_respects_master_switch(session: Session, plan, monkeypatch):
    from app.services import production_automation as pa

    calls: list = []
    monkeypatch.setattr(
        pa, "enqueue_plan_pipeline",
        lambda *a, **k: calls.append(a),
    )

    assert pa.maybe_trigger_plan_automation(session, plan.id, 1) is False
    assert not calls, "总开关关闭时不能投递任务"

    from app.services.production_automation_settings import save_automation_settings

    save_automation_settings(
        session, 1,
        {"enabled": True, "on_plan_saved": {"run_schedule": True}},
    )
    assert pa.maybe_trigger_plan_automation(session, plan.id, 1) is True
    assert calls and list(calls[0]) == [plan.id, 1, "plan_saved"]
